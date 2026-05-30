import os
from pathlib import Path


class FileParser:
    SUPPORTED_TYPES = {".pdf", ".docx", ".xlsx", ".txt", ".md", ".html", ".htm"}

    @classmethod
    def parse(cls, file_path: str) -> str:
        path = Path(file_path)
        suffix = path.suffix.lower()
        if suffix not in cls.SUPPORTED_TYPES:
            return f"[不支持的文件类型: {suffix}]"
        if suffix in {".txt", ".md"}:
            return cls._parse_text(file_path)
        elif suffix in {".html", ".htm"}:
            return cls._parse_html(file_path)
        elif suffix == ".docx":
            return cls._parse_docx(file_path)
        elif suffix == ".pdf":
            return cls._parse_pdf(file_path)
        elif suffix == ".xlsx":
            return cls._parse_xlsx(file_path)
        return "[解析失败: 未知类型]"

    @staticmethod
    def _parse_text(file_path: str) -> str:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()

    @staticmethod
    def _parse_html(file_path: str) -> str:
        try:
            from bs4 import BeautifulSoup
            with open(file_path, "r", encoding="utf-8") as f:
                soup = BeautifulSoup(f.read(), "html.parser")
                for tag in soup(["script", "style"]):
                    tag.decompose()
                return soup.get_text(separator="\n", strip=True)
        except ImportError:
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()

    @staticmethod
    def _parse_docx(file_path: str) -> str:
        try:
            from docx import Document
            doc = Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n".join(paragraphs)
        except ImportError:
            return "[解析失败: python-docx 未安装]"

    @staticmethod
    def _parse_pdf(file_path: str) -> str:
        try:
            import pdfplumber
            texts = []
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        texts.append(text)
            return "\n".join(texts)
        except ImportError:
            return "[解析失败: pdfplumber 未安装]"

    @staticmethod
    def _parse_xlsx(file_path: str) -> str:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            texts = []
            for sheet in wb.worksheets:
                texts.append(f"=== Sheet: {sheet.title} ===")
                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join(str(cell) for cell in row if cell is not None)
                    if row_text.strip():
                        texts.append(row_text)
            return "\n".join(texts)
        except ImportError:
            return "[解析失败: openpyxl 未安装]"
